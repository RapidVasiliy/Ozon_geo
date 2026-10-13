"""Excel-отчёт: вкладка = кластер, строки = SKU, столбцы = месяцы."""
from __future__ import annotations

import re

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.hyperlink import Hyperlink

from . import analytics

FONT = "Arial"
NOTES_SHEET = "Примечания"
METRICS = {"quantity": ("штуки", "#,##0;-#,##0;\"–\""), "revenue": ("выручка, руб.", "#,##0;-#,##0;\"–\"")}


def _num(v) -> int | float:
    """Целое для штук, число с копейками для выручки."""
    return int(v) if float(v).is_integer() else round(float(v), 2)


def _sheet_name(name: str, used: set[str]) -> str:
    """Имя вкладки: ≤31 символа, без `[]:*?/\\`, уникальное без учёта регистра."""
    base = re.sub(r"[\[\]:*?/\\]", " ", name).strip("' ")[:31] or "Лист"
    cand, k = base, 2
    while cand.lower() in used or cand == NOTES_SHEET:
        suffix = f" {k}"
        cand, k = base[: 31 - len(suffix)] + suffix, k + 1
    used.add(cand.lower())
    return cand


def cluster_month_workbook(df: pd.DataFrame, path: str, metric: str = "quantity", geo: str = "cluster",
                           source_note: str = "", formulas: bool = False) -> str:
    """Собрать xlsx. `df` — результат `analytics.prepare(...)`.

    Итоги по умолчанию — готовые числа: файл показывает значения в любом просмотрщике (формулы без
    сохранённых значений в превью выглядят пустыми, а пересчитать их нечем без Excel/LibreOffice).
    `formulas=True` — итоги формулами `=SUM(...)`, Excel посчитает их при открытии.
    """
    label, fmt = METRICS[metric]
    d = df.copy()
    d["month"] = d["created_at"].dt.strftime("%Y-%m")
    first, last = d["created_at"].min(), d["created_at"].max()
    months = list(pd.period_range(first.strftime("%Y-%m"), last.strftime("%Y-%m"), freq="M").strftime("%Y-%m"))
    partial = not last.is_month_end
    headers = [f"{m} (до {last:%d.%m})" if partial and m == months[-1] else m for m in months]
    names = d.groupby("sku")["name"].first()
    col = analytics.GEOS[geo]
    clusters = d.groupby(col)[metric].sum().sort_values(ascending=False).index.tolist()

    head_font, body_font = Font(name=FONT, bold=True, size=10), Font(name=FONT, size=10)
    head_fill = PatternFill("solid", fgColor="D9E1F2")
    total_fill = PatternFill("solid", fgColor="F2F2F2")
    thin = Side(style="thin", color="999999")

    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    totals: dict[str, float] = {}
    index: list[tuple[str, str, int, int, int]] = []  # (кластер, вкладка, SKU, строка итога, столбец итога)
    n_first, n_last = 3, 2 + len(months)  # столбцы месяцев: C..; «Итого» следом
    tot_col = n_last + 1

    for cl in clusters:
        sub = d[d[col] == cl]
        pv = (sub.pivot_table(index="sku", columns="month", values=metric, aggfunc="sum", fill_value=0)
                 .reindex(columns=months, fill_value=0))
        pv = pv.loc[pv.sum(axis=1).sort_values(ascending=False, kind="stable").index]
        title = _sheet_name(cl, used)
        ws = wb.create_sheet(title)
        for j, h in enumerate(["SKU", "Название", *headers, "Итого"], start=1):
            c = ws.cell(1, j, h)
            c.font, c.fill, c.border = head_font, head_fill, Border(bottom=thin)
            c.alignment = Alignment(horizontal="left" if j <= 2 else "center", vertical="center", wrap_text=True)
        for i, (sku, row) in enumerate(pv.iterrows(), start=2):
            ws.cell(i, 1, str(sku)).font = body_font  # текстом: чтобы Excel не превращал номер в число
            ws.cell(i, 2, str(names[sku])).font = body_font
            for j, v in enumerate(row.values, start=n_first):
                c = ws.cell(i, j, float(v) if metric == "revenue" else int(v))
                c.font, c.number_format = body_font, fmt
            c = ws.cell(i, tot_col, f"=SUM({L(n_first)}{i}:{L(n_last)}{i})" if formulas else _num(row.sum()))
            c.font, c.number_format = Font(name=FONT, size=10, bold=True), fmt
        end = len(pv) + 1
        tr = end + 1
        ws.cell(tr, 1, "Итого")
        for j in range(n_first, tot_col + 1):
            val = f"=SUM({L(j)}2:{L(j)}{end})" if formulas else _num(pv.iloc[:, j - n_first].sum() if j < tot_col else pv.values.sum())
            ws.cell(tr, j, val).number_format = fmt
        for j in range(1, tot_col + 1):
            c = ws.cell(tr, j)
            c.font, c.fill, c.border = head_font, total_fill, Border(top=thin)
        ws.freeze_panes = "C2"
        ws.column_dimensions["A"].width = 13
        ws.column_dimensions["B"].width = 62
        for j in range(n_first, tot_col + 1):
            ws.column_dimensions[L(j)].width = 12 if j < tot_col else 11
        ws.row_dimensions[1].height = 30
        ws.sheet_properties.tabColor = "4472C4"
        totals[cl] = _num(pv.values.sum())
        index.append((cl, title, len(pv), tr, tot_col))

    # --- «Примечания»: оглавление со ссылками и допущения -------------------------------------------
    ns = wb.create_sheet(NOTES_SHEET)
    ns["A1"] = f"Метрика: {label}. Кластер доставки × месяц × SKU"
    ns["A1"].font = Font(name=FONT, bold=True, size=12)
    for j, h in enumerate(["Кластер (вкладка)", "SKU с продажами", f"Итого, {label}"], start=1):
        c = ns.cell(3, j, h)
        c.font, c.fill = head_font, head_fill
    for i, (cl, title, n_sku, tr, tc) in enumerate(index, start=4):
        q = title.replace("'", "''")
        c = ns.cell(i, 1, title)
        c.hyperlink = Hyperlink(ref=c.coordinate, location=f"'{q}'!A1", display=title)  # внутренняя ссылка на вкладку
        c.font = Font(name=FONT, size=10, color="0563C1", underline="single")
        ns.cell(i, 2, n_sku).font = body_font
        c = ns.cell(i, 3, f"='{q}'!{L(tc)}{tr}" if formulas else totals[cl])
        c.font, c.number_format = Font(name=FONT, size=10, color="008000"), fmt  # ссылка на другой лист
    r = 4 + len(index)
    ns.cell(r, 1, "Всего")
    ns.cell(r, 3, f"=SUM(C4:C{r - 1})" if formulas else sum(totals.values())).number_format = fmt
    for j in (1, 2, 3):
        ns.cell(r, j).font, ns.cell(r, j).fill = head_font, total_fill
    notes = [
        "Допущения и пояснения:",
        f"• Период данных: {first:%d.%m.%Y} — {last:%d.%m.%Y}; всего отправлений: {d['posting_number'].nunique():,}.",
        "• Метрика «штуки» — сумма поля «Количество»; «выручка» — «Оплачено покупателем» × количество.",
        "• Отменённые отправления исключены; доставленные, в пути и ожидающие отгрузки/выдачи включены.",
        "• Месяц определяется по дате «Принят в обработку»: в выгрузке она в UTC (сверено с API), месяцы считаются по Москве (UTC+3).",
        "• Кластер — поле «Кластер доставки» (куда уехал заказ), а не кластер отгрузки.",
        "• На вкладке кластера — только SKU, у которых в этом кластере были продажи; порядок — по убыванию «Итого».",
        *( [f"• Последний месяц неполный: данные до {last:%d.%m.%Y}."] if partial else []),
        *( [f"• Источник: {source_note}"] if source_note else []),
    ]
    for k, t in enumerate(notes):
        c = ns.cell(r + 2 + k, 1, t)
        c.font = Font(name=FONT, size=10, bold=(k == 0))
    ns.column_dimensions["A"].width = 34
    ns.column_dimensions["B"].width = 16
    ns.column_dimensions["C"].width = 18
    wb.save(path)
    return path
