import re

import pytest
from openpyxl import load_workbook

from ozon_geo import analytics
from ozon_geo.demo import generate
from ozon_geo.store import Store
from ozon_geo.xlsx_report import NOTES_SHEET, _sheet_name, cluster_month_workbook


def _df():
    st = Store(":memory:")
    generate(st, days=150, seed=5)
    return analytics.prepare(st.sales_frame())


def test_sheet_name_rules():
    used: set[str] = set()
    assert _sheet_name("Москва, МО и Дальние регионы", used) == "Москва, МО и Дальние регионы"
    assert _sheet_name("a/b:c*d?[e]\\", used) == "a b c d  e"
    n1, n2 = _sheet_name("x" * 40, used), _sheet_name("x" * 40, used)
    assert len(n1) == 31 and len(n2) <= 31 and n1 != n2
    assert _sheet_name(NOTES_SHEET, used) != NOTES_SHEET  # имя служебной вкладки занято


def test_layout_values_and_formula_ranges(tmp_path):
    df = _df()
    path = cluster_month_workbook(df, str(tmp_path / "c.xlsx"))
    wb = load_workbook(path)
    clusters = df.groupby("cluster")["quantity"].sum().sort_values(ascending=False).index.tolist()
    assert wb.sheetnames[:-1] == clusters and wb.sheetnames[-1] == NOTES_SHEET  # вкладка = кластер, по убыванию объёма
    months = sorted(df["created_at"].dt.strftime("%Y-%m").unique())
    for cl in clusters:
        ws = wb[cl]
        header = [c.value for c in ws[1]]
        assert header[:2] == ["SKU", "Название"] and header[-1] == "Итого"
        assert [h[:7] for h in header[2:-1]] == months  # месяцы по столбцам
        sub = df[df["cluster"] == cl]
        skus = [ws.cell(r, 1).value for r in range(2, ws.max_row)]
        assert set(skus) == set(sub["sku"].astype(str))  # по строкам — SKU, только проданные в кластере
        for r in range(2, ws.max_row):
            row = [ws.cell(r, c).value for c in range(3, 3 + len(months))]
            exp = (sub[sub["sku"].astype(str) == ws.cell(r, 1).value]
                   .groupby(sub["created_at"].dt.strftime("%Y-%m"))["quantity"].sum())
            assert row == [int(exp.get(m, 0)) for m in months]
            assert ws.cell(r, 3 + len(months)).value == f"=SUM(C{r}:{chr(ord('C') + len(months) - 1)}{r})"
        last = ws.max_row
        assert ws.cell(last, 1).value == "Итого"
        assert re.fullmatch(rf"=SUM\(C2:C{last - 1}\)", ws.cell(last, 3).value)
        assert ws.freeze_panes == "C2"
    notes = wb[NOTES_SHEET]
    assert notes.cell(4, 1).value == clusters[0] and notes.cell(4, 1).hyperlink is not None
